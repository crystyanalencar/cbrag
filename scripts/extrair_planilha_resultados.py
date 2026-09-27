#!/usr/bin/env python
"""Extrai a "Planilha de Resultados" da Central de Downloads (.xlsx, formato
largo: uma linha por métrica, uma coluna por período) pra formato longo.

Lê todas as versões arquivadas por `baixar_ri_mziq.py` em
`data/ri_central/_planilhas/` e grava, em `data/cvm_estruturado/`:

- `planilha_resultados.json`: a versão mais recente, uma linha por
  (aba, grupo, métrica, período).
- `planilha_revisoes.json`: períodos cujo valor mudou entre versões
  (republicação), com o valor de cada versão. Vazio enquanto só há uma.

Antes de gravar, confere receita líquida e lucro líquido da aba DRE contra o
`dre.json` da CVM (trimestre isolado e ano fechado); divergência nos períodos
recentes barra a gravação. Estrutura da planilha e desenho: `docs/ingestao.md`.

Uso: `uv run python scripts/extrair_planilha_resultados.py [--sem-validar]`
"""
import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

from cbrag.dados_financeiros import _DIAS_ANO_FECHADO, _DIAS_TRIMESTRE

ROOT = Path(__file__).resolve().parents[1]
PLANILHAS_DIR = ROOT / "data/ri_central/_planilhas"
OUT_DIR = ROOT / "data/cvm_estruturado"
OUT_FILE = OUT_DIR / "planilha_resultados.json"
REVISOES_FILE = OUT_DIR / "planilha_revisoes.json"
DRE_CVM_FILE = OUT_DIR / "dre.json"

# Capa é só botões; Conciliação FC é uma matriz de um período (não é série).
ABAS_IGNORADAS = {"Capa", "Conciliação FC"}

# Tolerância na conferência com a CVM: a planilha está em milhões
# (arredondada), a CVM em milhares.
TOLERANCIA_MILHOES = 1.5
# Só os períodos recentes barram a gravação: números antigos são
# republicados pela companhia sem aviso (nota "republicados em 2014" na DRE).
PERIODOS_RECENTES = 8

_ARQUIVO_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})_([0-9a-f]{8})\.xlsx$")
_TRIMESTRE_RE = re.compile(r"^(\d)T(\d{2})$")
_ANO_RE = re.compile(r"^(\d{4})$")


def parse_periodo(valor) -> dict | None:
    """'2T26\\n2Q26' -> trimestre 2 de 2026; '2018' ou 2018 -> ano 2018."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)) and 1990 <= int(valor) <= 2100 and int(valor) == valor:
        return {"periodo": str(int(valor)), "tipo_periodo": "ano", "ano": int(valor), "trimestre": None}
    texto = str(valor).strip().split("\n")[0].strip()
    m = _TRIMESTRE_RE.match(texto)
    if m:
        return {"periodo": texto, "tipo_periodo": "trimestre", "ano": 2000 + int(m.group(2)), "trimestre": int(m.group(1))}
    m = _ANO_RE.match(texto)
    if m:
        return {"periodo": texto, "tipo_periodo": "ano", "ano": int(texto), "trimestre": None}
    return None


def achar_cabecalho(ws) -> tuple[int, dict[int, dict]] | None:
    """Primeira linha (das 12 primeiras) com pelo menos 3 células de período.
    Devolve (número da linha, {índice da coluna: período})."""
    for numero in range(1, 13):
        periodos = {}
        for i, celula in enumerate(ws[numero]):
            p = parse_periodo(celula.value)
            if p:
                periodos[i] = p
        if len(periodos) >= 3:
            return numero, periodos
    return None


def inferir_unidade(aba: str, rotulo_pt: str, cabecalho_a: str, formato: str) -> str:
    rotulo = rotulo_pt.lower()
    # "%" solto no rótulo não basta: "Deságio de 15%" é um R$ milhões.
    if "%" in formato or "(%)" in rotulo or rotulo.rstrip().endswith("%"):
        return "%"
    if formato.rstrip().endswith("x"):  # múltiplo, ex.: covenant "0.0\x"
        return "x"
    if "mil m²" in rotulo or "mil m2" in rotulo:
        return "mil m²"
    if aba == "Lojas":
        return "unidade"
    # Demais abas são financeiras, em R$ milhões (o cabeçalho ou o rótulo
    # dizem "milhões de R$"/"R$ milhões"); GMV não diz, mas também é milhões.
    return "R$ milhões"


def _numero(valor):
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        return None
    return valor


def extrair_versao(caminho: Path, versao: str) -> list[dict]:
    """Uma linha por (aba, grupo, métrica, período) com valor numérico."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    registros = []
    for ws in wb.worksheets:
        aba = ws.title.split(" | ")[0].strip()
        if aba in ABAS_IGNORADAS:
            continue
        achado = achar_cabecalho(ws)
        if achado is None:
            print(f"AVISO aba '{ws.title}': cabeçalho de período não encontrado, ignorada")
            continue
        linha_cab, periodos = achado
        ultima_coluna = max(periodos)
        cabecalho_a = str(ws.cell(row=linha_cab, column=1).value or "")
        # A hierarquia está no recuo (`alignment.indent`) da coluna A, não em
        # linha vazia: em Lojas "Casas Bahia" tem número (total de lojas) e
        # "Abertas/Fechadas" vêm recuadas embaixo dela.
        pais: dict[int, str] = {}
        vistas: dict[tuple, int] = {}
        for row in ws.iter_rows(min_row=linha_cab + 1, max_col=ultima_coluna + 1):
            rotulo_pt = str(row[0].value).strip() if row[0].value is not None else ""
            rotulo_en = str(row[1].value).strip() if len(row) > 1 and row[1].value is not None else ""
            if not rotulo_pt and not rotulo_en:
                continue
            recuo = int(row[0].alignment.indent or 0)
            for nivel in [n for n in pais if n >= recuo]:
                del pais[nivel]
            grupo = " > ".join(pais[n] for n in sorted(pais))
            pais[recuo] = rotulo_pt or rotulo_en
            valores = {i: _numero(row[i].value) for i in periodos if i < len(row)}
            if all(v is None for v in valores.values()):
                # Sem número em nenhum período: cabeçalho de grupo (ou
                # rodapé, que não gera registro nenhum).
                continue
            chave = (aba, grupo, rotulo_pt)
            vistas[chave] = vistas.get(chave, 0) + 1
            formato = row[ultima_coluna].number_format if ultima_coluna < len(row) else ""
            unidade = inferir_unidade(aba, rotulo_pt, cabecalho_a, formato)
            for i, valor in valores.items():
                if valor is None:
                    continue
                registros.append({
                    "aba": aba,
                    "grupo": grupo,
                    "metrica_pt": rotulo_pt,
                    "metrica_en": rotulo_en,
                    "ocorrencia": vistas[chave],
                    **periodos[i],
                    "valor": valor,
                    "unidade": unidade,
                    "versao": versao,
                })
    return registros


def versoes_arquivadas() -> list[tuple[str, Path]]:
    achados = []
    for caminho in PLANILHAS_DIR.glob("*.xlsx"):
        m = _ARQUIVO_RE.match(caminho.name)
        if m:
            achados.append((m.group(1), caminho))
    return sorted(achados, key=lambda par: (par[0], par[1].name))


def _chave(r: dict) -> tuple:
    return (r["aba"], r["grupo"], r["metrica_pt"], r["ocorrencia"], r["periodo"])


def calcular_revisoes(por_versao: dict[str, list[dict]]) -> list[dict]:
    """Períodos presentes em mais de uma versão cujo valor mudou."""
    valores: dict[tuple, dict[str, float]] = {}
    for versao, registros in por_versao.items():
        for r in registros:
            valores.setdefault(_chave(r), {})[versao] = r["valor"]
    revisoes = []
    for chave, por in valores.items():
        if len(por) < 2:
            continue
        vs = list(por.values())
        if max(vs) - min(vs) > 1e-9:
            aba, grupo, metrica, ocorrencia, periodo = chave
            revisoes.append({
                "aba": aba, "grupo": grupo, "metrica_pt": metrica, "ocorrencia": ocorrencia,
                "periodo": periodo, "valores": [{"versao": v, "valor": x} for v, x in sorted(por.items())],
            })
    return revisoes


def conferir_com_cvm(registros: list[dict]) -> tuple[list[str], list[str]]:
    """Receita líquida (3.01) e lucro líquido (3.11) da aba DRE contra a
    DRE da CVM. Devolve (divergências em períodos recentes, demais)."""
    if not DRE_CVM_FILE.exists():
        return [], ["dre.json da CVM não encontrado: conferência pulada"]
    cvm = json.loads(DRE_CVM_FILE.read_text(encoding="utf-8"))
    contas = {"3.01": "Receita Líquida", "3.11": "Lucro Líquido"}
    esperado: dict[tuple, float] = {}
    for linha in cvm:
        if linha.get("CD_CONTA") not in contas or "LTIMO" not in linha.get("ORDEM_EXERC", "").upper():
            continue
        ini = date.fromisoformat(linha["DT_INI_EXERC"])
        fim = date.fromisoformat(linha["DT_FIM_EXERC"])
        dias = (fim - ini).days
        if _DIAS_TRIMESTRE[0] <= dias <= _DIAS_TRIMESTRE[1]:
            periodo = f"{(fim.month - 1) // 3 + 1}T{fim.year % 100:02d}"
        elif _DIAS_ANO_FECHADO[0] <= dias <= _DIAS_ANO_FECHADO[1]:
            periodo = str(fim.year)
        else:
            continue
        esperado[(contas[linha["CD_CONTA"]], periodo)] = float(linha["VL_CONTA"]) / 1000

    dre = [r for r in registros if r["aba"] == "DRE"]
    achados = {}
    for r in dre:
        for nome in contas.values():
            if r["metrica_pt"].startswith(nome) and "%" not in r["unidade"] and r["ocorrencia"] == 1:
                achados[(nome, r["periodo"])] = r
    recentes = sorted({p for _, p in achados if _TRIMESTRE_RE.match(p)}, key=lambda p: (p[-2:], p[0]))[-PERIODOS_RECENTES:]
    recentes += [p for p in {p for _, p in achados if _ANO_RE.match(p)} if int(p) >= date.today().year - 3]

    barram, outras = [], []
    comparados = 0
    for (nome, periodo), r in sorted(achados.items(), key=lambda kv: (kv[0][1][-2:], kv[0][1], kv[0][0])):
        if (nome, periodo) not in esperado:
            continue
        comparados += 1
        diff = abs(r["valor"] - esperado[(nome, periodo)])
        if diff > TOLERANCIA_MILHOES:
            texto = f"{nome} {periodo}: planilha {r['valor']:.0f} x CVM {esperado[(nome, periodo)]:.0f} (diferença {diff:.0f} milhões)"
            (barram if periodo in recentes else outras).append(texto)
    outras.insert(0, f"{comparados} valores conferidos com a CVM")
    return barram, outras


def _gravar_atomico(destino: Path, texto: str) -> None:
    """Escreve num temporário e renomeia: o chat lê estes arquivos enquanto a
    ingestão roda e nunca deve pegar um arquivo pela metade."""
    temporario = destino.with_name(destino.name + ".tmp")
    temporario.write_text(texto, encoding="utf-8")
    temporario.replace(destino)


def main(validar: bool = True) -> dict:
    versoes = versoes_arquivadas()
    if not versoes:
        print(f"Nenhuma planilha em {PLANILHAS_DIR}. Nada a fazer.")
        return {"planilha_versoes": 0}

    por_versao = {}
    for versao, caminho in versoes:
        # Duas cópias com a mesma data de entrega (mesmo dia, conteúdo
        # diferente): a última, em ordem de nome, prevalece.
        por_versao[versao] = extrair_versao(caminho, versao)
        print(f"{caminho.name}: {len(por_versao[versao])} valores")

    mais_recente = max(por_versao)
    registros = por_versao[mais_recente]

    if validar:
        barram, outras = conferir_com_cvm(registros)
        for texto in outras:
            print(f"  {texto}")
        if barram:
            print("DIVERGÊNCIA com a CVM em período recente, nada gravado:", file=sys.stderr)
            for texto in barram:
                print(f"  {texto}", file=sys.stderr)
            raise SystemExit(1)

    revisoes = calcular_revisoes(por_versao)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    # Uma linha por registro: dá diff legível e o arquivo (milhares de
    # linhas) não vira uma linha só.
    corpo = ",\n".join(json.dumps(r, ensure_ascii=False) for r in registros)
    _gravar_atomico(OUT_FILE, f"[\n{corpo}\n]\n")
    _gravar_atomico(REVISOES_FILE, json.dumps(revisoes, indent=2, ensure_ascii=False))
    print(f"{len(registros)} valores da versão {mais_recente} em {OUT_FILE}; {len(revisoes)} revisões entre {len(versoes)} versão(ões)")
    return {"planilha_versoes": len(versoes), "planilha_valores": len(registros), "planilha_revisoes": len(revisoes)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sem-validar", action="store_true", help="não confere com a DRE da CVM")
    args = parser.parse_args()
    main(validar=not args.sem_validar)
